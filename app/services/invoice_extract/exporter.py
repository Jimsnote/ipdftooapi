"""台账导出：CSV（UTF-8 BOM，否则 Excel 打开中文乱码）/ xlsx（openpyxl）。

导出以前端表格当前态为准（用户手动修正/补录的行原样落表），后端只做格式化。
"""
import csv
import io
from datetime import datetime
from typing import List, Optional

from app.models.invoice_extract import ExportRowRequest

# (列头, 字段名, 是否金额列)
COLUMNS = [
    ("来源文件名", "source_file", False),
    ("发票类型", "invoice_type", False),
    ("发票号码", "invoice_number", False),
    ("开票日期", "issue_date", False),
    ("购买方名称", "buyer_name", False),
    ("购买方税号", "buyer_tax_id", False),
    ("销售方名称", "seller_name", False),
    ("销售方税号", "seller_tax_id", False),
    ("不含税金额", "amount_without_tax", True),
    ("税额", "tax_amount", True),
    ("价税合计", "total_with_tax", True),
    ("明细行数", "item_count", False),
    ("备注", "remark", False),
    ("校验提示", "warnings", False),
]

MONEY_FORMAT = "0.00"


def _cell(row: ExportRowRequest, key: str):
    if key == "warnings":
        return "；".join(row.warnings or []) or None
    value = getattr(row, key)
    if isinstance(value, str) and not value.strip():
        return None
    return value


def _as_number(value) -> Optional[float]:
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.replace(",", "").replace("¥", "").replace("￥", "").strip()
        try:
            return float(s)
        except ValueError:
            return None
    return None


def export_csv(rows: List[ExportRowRequest]) -> bytes:
    """CSV 字节串。首字符为 UTF-8 BOM（Excel 兼容关键细节）。"""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([h for h, _, _ in COLUMNS])
    for row in rows:
        writer.writerow([("" if _cell(row, k) is None else _cell(row, k)) for _, k, _ in COLUMNS])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def export_xlsx(rows: List[ExportRowRequest]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "发票台账"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")

    for col, (header, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, start=2):
        for col, (_, key, is_money) in enumerate(COLUMNS, start=1):
            value = _cell(row, key)
            cell = ws.cell(row=r, column=col)
            if value is None:
                continue
            if is_money:
                num = _as_number(value)
                if num is not None:
                    cell.value = num
                    cell.number_format = MONEY_FORMAT
                else:
                    cell.value = str(value)
            elif key == "item_count":
                num = _as_number(value)
                cell.value = int(num) if num is not None else str(value)
            else:
                cell.value = str(value)

    widths = [24, 24, 24, 12, 30, 22, 30, 22, 14, 12, 14, 10, 32, 32]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows) + 1)}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_filename(fmt: str) -> str:
    ext = "xlsx" if fmt == "xlsx" else "csv"
    return f"发票台账_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"
