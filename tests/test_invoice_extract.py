"""发票信息批量提取测试：三适配器 + 导出器 + 路由接口。

- XML 夹具内联构造（schema 与真实样本一致，见 docs/xml-invoice/）；
- PDF 夹具用 fitz 按官方横版几何合成文字层；扫描件用无文字页；
- OFD 夹具在内存里构造最小 ZIP（CustomTag 精确路径 / 无标签回退路径各一）。
"""
import io
import zipfile
import xml.etree.ElementTree as defused_et  # 仅测试夹具构造用标准库

import fitz
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services.invoice_extract import (
    InvoiceExtractError,
    extract_bytes,
    export_csv,
    export_xlsx,
)

client = TestClient(app)

# ---------------------------------------------------------------------------
# XML 夹具
# ---------------------------------------------------------------------------

XML_TMPL = """<?xml version="1.0" encoding="utf-8"?>
<EInvoice>
  <Header>
    <EIid>{number}</EIid>
    <InherentLabel>
      <GeneralOrSpecialVAT><LabelName>{vat}</LabelName></GeneralOrSpecialVAT>
      <InIssuType><LabelCode>{blue}</LabelCode></InIssuType>
    </InherentLabel>
  </Header>
  <TaxSupervisionInfo>
    <InvoiceNumber>{number}</InvoiceNumber>
    <IssueTime>{date}</IssueTime>
    <TaxBureauName>国家税务总局北京市税务局</TaxBureauName>
  </TaxSupervisionInfo>
  <EInvoiceData>
    <BuyerInformation>
      <BuyerName>{buyer_name}</BuyerName><BuyerIdNum>{buyer_tax}</BuyerIdNum>
    </BuyerInformation>
    <SellerInformation>
      <SellerName>{seller_name}</SellerName><SellerIdNum>{seller_tax}</SellerIdNum>
    </SellerInformation>
    <BasicInformation>
      <TotalAmWithoutTax>{without}</TotalAmWithoutTax>
      <TotalTaxAm>{tax}</TotalTaxAm>
      <TotalTax-includedAmount>{total}</TotalTax-includedAmount>
      <TotalTax-includedAmountInChinese>贰佰柒拾圆整</TotalTax-includedAmountInChinese>
      <Drawer>张三</Drawer>
    </BasicInformation>
    <IssuItemInformation><ItemName>*餐饮服务*餐费</ItemName></IssuItemInformation>
    <Remark>{remark}</Remark>
  </EInvoiceData>
</EInvoice>
"""


def _xml(**kw) -> bytes:
    defaults = dict(
        number="26112000001182262471", vat="普通发票", blue="Y", date="2026-03-26",
        buyer_name="北京创信卓远信息技术有限责任公司", buyer_tax="9111010859062383XH",
        seller_name="方叔叔餐饮管理（北京）有限公司", seller_tax="91110105MA0084MW37",
        without="254.72", tax="15.28", total="270.00", remark="",
    )
    defaults.update(kw)
    return XML_TMPL.format(**defaults).encode("utf-8")


class TestXmlAdapter:
    def test_standard_fields(self):
        rec = extract_bytes(_xml(), "a.xml", ".xml")
        assert rec.invoice_number == "26112000001182262471"
        assert rec.issue_date == "2026-03-26"
        assert rec.invoice_type == "电子发票（普通发票）"
        assert rec.buyer_name == "北京创信卓远信息技术有限责任公司"
        assert rec.seller_tax_id == "91110105MA0084MW37"
        assert (rec.amount_without_tax, rec.tax_amount, rec.total_with_tax) == ("254.72", "15.28", "270.00")
        assert rec.item_count == 1

    def test_red_letter_warning(self):
        rec = extract_bytes(_xml(blue="N"), "b.xml", ".xml")
        assert any("红字" in w for w in rec.warnings)

    def test_zh_cn_date_normalized(self):
        rec = extract_bytes(_xml(date="2026年03月26日"), "c.xml", ".xml")
        assert rec.issue_date == "2026-03-26"

    def test_not_einvoice_rejected(self):
        with pytest.raises(InvoiceExtractError):
            extract_bytes(b"<Order><id>1</id></Order>", "d.xml", ".xml")

    def test_malformed_xml_rejected(self):
        with pytest.raises(InvoiceExtractError):
            extract_bytes(b"<EInvoice><unclosed>", "e.xml", ".xml")

    def test_missing_number_rejected(self):
        bad = _xml().replace(b"<InvoiceNumber>26112000001182262471</InvoiceNumber>", b"")
        bad = bad.replace(b"<EIid>26112000001182262471</EIid>", b"")
        with pytest.raises(InvoiceExtractError):
            extract_bytes(bad, "f.xml", ".xml")


# ---------------------------------------------------------------------------
# PDF 夹具（fitz 合成官方横版文字层）
# ---------------------------------------------------------------------------

def _make_pdf_bytes(with_text: bool = True) -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=595.28, height=396.85)
    if with_text:
        def put(x, y, s):
            # 中文用内置 china-s；纯 ASCII 用 helv（china-s 的 ASCII 是全角宽，会溢出裁剪）
            font = "helv" if s.isascii() else "china-s"
            page.insert_text((x, y), s, fontsize=9, fontname=font)

        put(161.7, 40, "电子发票（增值税专用发票）")
        put(438.0, 39.4, "发票号码：")
        put(484.0, 39.4, "26112000003559581871")
        put(438.0, 56.6, "开票日期：")
        put(484.0, 56.6, "2026年08月25日")
        put(32.7, 103.5, "名称：")
        put(57.0, 103.5, "中国人民财产保险股份有限公司")
        put(317.6, 103.5, "名称：")
        put(341.0, 103.5, "北京创信卓远信息技术有限公司")
        put(153.1, 132.6, "91100000710931483R")
        put(437.9, 132.6, "9111010859062383XH")
        put(13.8, 167.9, "*软件服务*技术服务费")
        put(394.6, 268.5, "695086.79")   # 官方版式 ¥ 与数字分字体，这里仅放数字
        put(546.5, 268.5, "41705.21")
        put(406.8, 287.7, "（小写）")
        put(447.4, 287.7, "¥736792.00")  # 故意与标签同字体相邻，覆盖"合并 span 兜底"路径
        put(34.0, 318.0, "合同编号HT2026-0825")
    data = doc.tobytes()
    doc.close()
    return data


class TestPdfAdapter:
    def test_text_layer_invoice(self):
        rec = extract_bytes(_make_pdf_bytes(), "inv.pdf", ".pdf")
        assert rec.source_format == "pdf"
        assert rec.invoice_number == "26112000003559581871"
        assert rec.issue_date == "2026-08-25"
        assert rec.invoice_type == "电子发票（增值税专用发票）"
        assert rec.buyer_name == "中国人民财产保险股份有限公司"
        assert rec.seller_name == "北京创信卓远信息技术有限公司"
        assert (rec.amount_without_tax, rec.tax_amount, rec.total_with_tax) == (
            "695086.79", "41705.21", "736792.00",
        )
        assert rec.item_count == 1
        assert rec.remark and "合同编号" in rec.remark

    def test_scanned_pdf_rejected(self):
        with pytest.raises(InvoiceExtractError) as ei:
            extract_bytes(_make_pdf_bytes(with_text=False), "scan.pdf", ".pdf")
        assert "文字层" in str(ei.value)


# ---------------------------------------------------------------------------
# OFD 夹具（内存构造最小 ZIP）
# ---------------------------------------------------------------------------

_OFD_NS = "http://www.ofdspec.org/2016"


def _page_xml(entries):
    """entries: [(id, x, y, text)] → Content.xml"""
    objs = "".join(
        f'<ofd:TextObject ID="{tid}" Boundary="{x} {y} 60 5">'
        f"<ofd:TextCode>{text}</ofd:TextCode></ofd:TextObject>"
        for tid, x, y, text in entries
    )
    return (
        f'<ofd:Page xmlns:ofd="{_OFD_NS}"><ofd:Area><ofd:PhysicalBox>0 0 210 140</ofd:PhysicalBox>'
        f"</ofd:Area><ofd:Content><ofd:Layer>{objs}</ofd:Layer></ofd:Content></ofd:Page>"
    ).encode("utf-8")


def _custom_tag_xml(refs) -> bytes:
    body = "".join(
        f"<ofd:{tag}>{''.join(f'<ofd:ObjectRef>{rid}</ofd:ObjectRef>' for rid in ids)}</ofd:{tag}>"
        for tag, ids in refs
    )
    return f'<ofd:root xmlns:ofd="{_OFD_NS}">{body}</ofd:root>'.encode("utf-8")


def _zip_ofd(pages: dict, tags: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("OFD.xml", "<ofd:OFDDoc xmlns:ofd='%s'/>" % _OFD_NS)
        for name, content in pages.items():
            zf.writestr(name, content)
        for name, content in tags.items():
            zf.writestr(name, content)
    return buf.getvalue()


PAGE_ENTRIES = [
    ("9001", 170, 10.3, "26112233445566778899"),
    ("9002", 170, 16.4, "2026年06月03日"),
    ("9003", 20, 33.8, "海港人寿保险股份有限公司"),
    ("9004", 54, 43.3, "91440300MACNKMUC8X"),
    ("9005", 120, 33.8, "北京创信卓远信息技术有限责任公司"),
    ("9006", 154, 43.3, "9111010859062383XH"),
    ("9007", 135.5, 93.8, "¥11320.75"),
    ("9008", 188.5, 93.8, "¥679.25"),
    ("9009", 179.5, 100.0, "¥12000.00"),
]
TAG_REFS = [
    ("InvoiceNo", ["9001"]),
    ("IssueDate", ["9002"]),
    ("BuyerName", ["9003"]),
    ("BuyerTaxID", ["9004"]),
    ("SellerName", ["9005"]),
    ("SellerTaxID", ["9006"]),
    ("TaxExclusiveTotalAmount", ["x-yuan", "9007"]),  # 首个引用是 ¥ 符号
    ("TaxTotalAmount", ["9008"]),
    ("TaxInclusiveTotalAmount", ["9009"]),
]


class TestOfdAdapter:
    def test_custom_tag_precise_path(self):
        data = _zip_ofd(
            {"Doc_0/Pages/Page_0/Content.xml": _page_xml(PAGE_ENTRIES)},
            {"Doc_0/Tags/CustomTag.xml": _custom_tag_xml(TAG_REFS)},
        )
        rec = extract_bytes(data, "a.ofd", ".ofd")
        assert rec.source_format == "ofd"
        assert rec.invoice_number == "26112233445566778899"
        assert rec.buyer_name == "海港人寿保险股份有限公司"
        assert rec.seller_tax_id == "9111010859062383XH"
        # 金额字段引用里的 "¥" 符号应被跳过
        assert rec.amount_without_tax == "11320.75"
        assert rec.tax_amount == "679.25"
        assert rec.total_with_tax == "12000.00"

    def test_fallback_zone_without_custom_tag(self):
        data = _zip_ofd(
            {
                "Doc_0/Pages/Page_0/Content.xml": _page_xml(PAGE_ENTRIES),
                "Doc_0/Tpls/Tpl_0/Content.xml": _page_xml(
                    [("9100", 100, 11, "电子发票（普通发票）")]
                ),
            },
            {},
        )
        rec = extract_bytes(data, "b.ofd", ".ofd")
        assert rec.invoice_number == "26112233445566778899"
        assert rec.issue_date == "2026-06-03"
        assert rec.invoice_type == "电子发票（普通发票）"
        assert rec.buyer_name == "海港人寿保险股份有限公司"
        assert rec.seller_name == "北京创信卓远信息技术有限责任公司"
        assert rec.amount_without_tax == "11320.75"

    def test_unrelated_zip_rejected(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("hello.txt", "hi")
        with pytest.raises(InvoiceExtractError):
            extract_bytes(buf.getvalue(), "c.ofd", ".ofd")


# ---------------------------------------------------------------------------
# 导出器
# ---------------------------------------------------------------------------

def _row():
    from app.models.invoice_extract import ExportRowRequest

    return ExportRowRequest(
        source_file="a.xml", source_format="xml", invoice_type="电子发票（普通发票）",
        invoice_number="26112000001182262471", issue_date="2026-03-26",
        buyer_name="北京创信卓远信息技术有限责任公司", buyer_tax_id="9111010859062383XH",
        seller_name="方叔叔餐饮管理（北京）有限公司", seller_tax_id="91110105MA0084MW37",
        amount_without_tax="254.72", tax_amount="15.28", total_with_tax="270.00",
        item_count=1, remark=None, warnings=["测试提示"],
    )


class TestExporter:
    def test_csv_has_utf8_bom(self):
        data = export_csv([_row()])
        assert data.startswith(b"\xef\xbb\xbf")
        text = data.decode("utf-8-sig")
        assert "来源文件名" in text and "校验提示" in text
        assert "测试提示" in text

    def test_xlsx_roundtrip(self):
        from openpyxl import load_workbook

        data = export_xlsx([_row()])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "来源文件名"
        assert ws.cell(row=2, column=9).value == 254.72  # 金额列写成数字
        assert isinstance(ws.cell(row=2, column=9).value, float)


# ---------------------------------------------------------------------------
# 路由接口
# ---------------------------------------------------------------------------

class TestApiRoutes:
    def test_analyze_mixed_batch(self):
        files = [
            ("files", ("ok.xml", _xml(), "application/xml")),
            ("files", ("bad.txt", b"not an invoice", "text/plain")),
            ("files", ("scan.pdf", _make_pdf_bytes(with_text=False), "application/pdf")),
        ]
        resp = client.post("/api/v1/invoice/extract-analyze", files=files)
        assert resp.status_code == 200
        body = resp.json()
        assert body["success_count"] == 1
        assert body["failure_count"] == 2
        reasons = " ".join(f["reason"] for f in body["failures"])
        assert "不支持的文件类型" in reasons
        assert "文字层" in reasons
        rec = body["records"][0]
        assert rec["invoice_number"] == "26112000001182262471"
        assert rec["total_with_tax"] == "270.00"

    def test_duplicate_numbers_warned(self):
        files = [
            ("files", (f"{i}.xml", _xml(), "application/xml")) for i in range(2)
        ]
        resp = client.post("/api/v1/invoice/extract-analyze", files=files)
        body = resp.json()
        assert body["success_count"] == 2
        for rec in body["records"]:
            assert any("出现 2 次" in w for w in rec["warnings"])

    def test_too_many_files_400(self):
        files = [("files", (f"{i}.xml", _xml(), "application/xml")) for i in range(51)]
        resp = client.post("/api/v1/invoice/extract-analyze", files=files)
        assert resp.status_code == 400

    def test_export_csv_roundtrip(self):
        rows = [_row().model_dump()]
        resp = client.post(
            "/api/v1/invoice/extract-export", json={"format": "csv", "rows": rows}
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["filename"].endswith(".csv") and body["count"] == 1
        dl = client.get(body["download_url"])
        assert dl.status_code == 200
        assert dl.content.startswith(b"\xef\xbb\xbf")

    def test_export_xlsx_roundtrip(self):
        from openpyxl import load_workbook

        rows = [_row().model_dump()]
        resp = client.post(
            "/api/v1/invoice/extract-export", json={"format": "xlsx", "rows": rows}
        )
        assert resp.status_code == 200
        body = resp.json()
        dl = client.get(body["download_url"])
        assert dl.status_code == 200
        wb = load_workbook(io.BytesIO(dl.content))
        assert wb.active.cell(row=1, column=3).value == "发票号码"

    def test_export_bad_format_400(self):
        resp = client.post(
            "/api/v1/invoice/extract-export", json={"format": "pdf", "rows": [_row().model_dump()]}
        )
        assert resp.status_code == 400
